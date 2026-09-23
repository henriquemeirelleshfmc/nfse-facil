"""Testes dos relatórios gerados a partir do arquivo interno de XMLs."""

from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from nfse_facil.domain.models import Empresa, PeriodoConsulta, PreferenciasDownload
from nfse_facil.services.documentos_adn import OrganizadorDocumentosADN
from nfse_facil.services.relatorios import GeradorRelatoriosService


def _preparar_xml(empresa: Empresa) -> None:
    pasta = (
        OrganizadorDocumentosADN.pasta_da_empresa(empresa)
        / OrganizadorDocumentosADN.PASTA_ARQUIVO_INTERNO
        / "2026" / "08" / "Emitidas"
    )
    pasta.mkdir(parents=True)
    pasta.joinpath("NFSE_CHAVE123.xml").write_text(
        """<NFSe><infNFSe Id="NFSCHAVE123"><nNFSe>789</nNFSe>
        <dhEmi>2026-08-15T10:00:00-04:00</dhEmi><xLocEmi>Manaus</xLocEmi>
        <xLocIncid>Manaus</xLocIncid>
        <valores><vBC>1500.50</vBC><pAliqAplic>5.0</pAliqAplic><vISSQN>75.03</vISSQN>
        <vTotalRet>30.00</vTotalRet><vLiq>1470.50</vLiq></valores>
        <DPS><infDPS><dCompet>2026-08-01</dCompet><serie>900</serie><nDPS>456</nDPS>
        <prest><CNPJ>00000000000191</CNPJ><xNome>Empresa Modelo</xNome></prest>
        <toma><CNPJ>11222333000181</CNPJ><xNome>Cliente Teste</xNome></toma>
        <serv><cServ><cTribNac>010101</cTribNac><xDescServ>Serviço de teste</xDescServ></cServ></serv>
        <valores><vServPrest><vServ>1500.50</vServ></vServPrest><trib><tribMun><tpRetISSQN>1</tpRetISSQN></tribMun>
        <tribFed><piscofins><vPis>10.00</vPis><vCofins>20.00</vCofins></piscofins></tribFed></trib></valores>
        </infDPS></DPS></infNFSe></NFSe>""",
        encoding="utf-8",
    )


def test_gera_excel_e_pdf_validos_com_dados_do_periodo(tmp_path: Path) -> None:
    empresa = Empresa("Empresa Modelo", "00000000000191", tmp_path / "docs")
    _preparar_xml(empresa)
    preferencias = PreferenciasDownload(
        baixar_emitidas=True, baixar_recebidas=True, gerar_excel=True, gerar_pdf=True
    )

    caminhos = GeradorRelatoriosService().gerar(
        empresa,
        PeriodoConsulta.personalizado(date(2026, 8, 1), date(2026, 8, 31)),
        preferencias,
    )

    assert {caminho.suffix for caminho in caminhos} == {".xlsx", ".pdf"}
    excel = next(caminho for caminho in caminhos if caminho.suffix == ".xlsx")
    pdf = next(caminho for caminho in caminhos if caminho.suffix == ".pdf")

    wb = load_workbook(excel, data_only=False)
    assert wb.sheetnames == ["Resumo emitidas", "Notas emitidas"]
    resumo = wb["Resumo emitidas"]
    ws = wb["Notas emitidas"]
    assert resumo["A2"].value == "Resumo fiscal - Notas emitidas"
    assert resumo["B10"].value == 1
    contatos = [cell for row in resumo.iter_rows() for cell in row if cell.hyperlink]
    assert any(cell.hyperlink.target == "https://www.instagram.com/henrique.meirelles_/" for cell in contatos)
    assert ws["A2"].value.date() == date(2026, 8, 15)
    assert ws["D2"].value == "789"
    assert ws["W2"].value == 1500.5
    assert ws["Y2"].value == 1500.5
    assert ws["AA2"].value == 75.03
    assert ws["AG2"].value == 30
    assert ws["AH2"].value == 1470.5
    assert ws.freeze_panes == "A2"
    assert len(ws.tables) == 1

    leitor = PdfReader(pdf)
    assert len(leitor.pages) >= 2
    texto = "\n".join(pagina.extract_text() or "" for pagina in leitor.pages)
    assert "Resumo Fiscal - Notas Emitidas" in texto
    assert "Empresa Modelo" in texto
    assert "Resumo Fiscal - Notas Emitidas" in texto
    assert "Detalhamento - Notas Emitidas" in texto
    assert "Total retido informado" in texto
    assert "@henrique.meirelles_" in texto
    assert "Tipo" not in texto


def test_respeita_opcoes_de_relatorio_desmarcadas(tmp_path: Path) -> None:
    empresa = Empresa("Empresa Modelo", "00000000000191", tmp_path / "docs")
    caminhos = GeradorRelatoriosService().gerar(
        empresa, PeriodoConsulta.este_mes(),
        PreferenciasDownload(gerar_excel=False, gerar_pdf=False),
    )
    assert caminhos == ()
    assert not OrganizadorDocumentosADN.pasta_da_empresa(empresa).exists()
