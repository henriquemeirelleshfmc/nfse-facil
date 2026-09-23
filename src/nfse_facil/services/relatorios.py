"""Relatórios fiscais completos em Excel e PDF a partir dos XMLs do ADN."""

from dataclasses import dataclass
from datetime import date
from pathlib import Path
import os
import re
import tempfile
from xml.etree import ElementTree

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table as ExcelTable, TableStyleInfo
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import HRFlowable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table as PDFTable, TableStyle

from nfse_facil.config.contact import SUPPORT_HANDLE, SUPPORT_URL
from nfse_facil.domain.document_models import DirecaoDocumento
from nfse_facil.domain.exceptions import DocumentoFiscalGravacaoError
from nfse_facil.domain.models import Empresa, PeriodoConsulta, PreferenciasDownload
from nfse_facil.services.documentos_adn import OrganizadorDocumentosADN


@dataclass(frozen=True)
class LinhaRelatorio:
    data_emissao: date
    competencia: date | None
    direcao: str
    numero_nfse: str
    serie_dps: str
    numero_dps: str
    chave: str
    municipio_emissao: str
    municipio_incidencia: str
    codigo_servico: str
    descricao_servico: str
    codigo_nbs: str
    descricao_nbs: str
    prestador_cnpj: str
    prestador_nome: str
    tomador_cnpj: str
    tomador_nome: str
    intermediario_cnpj: str
    intermediario_nome: str
    simples_nacional: str
    regime_especial: str
    tipo_retencao_iss: str
    valor_servico: float | None
    desconto: float | None
    base_iss: float | None
    aliquota_iss: float | None
    valor_iss: float | None
    valor_pis: float | None
    valor_cofins: float | None
    valor_irrf: float | None
    valor_csll: float | None
    valor_inss: float | None
    total_retido: float | None
    valor_liquido: float | None
    tributos_federais: float | None
    tributos_estaduais: float | None
    tributos_municipais: float | None
    total_tributos: float | None
    base_ibs_cbs: float | None
    ibs_uf: float | None
    ibs_municipal: float | None
    ibs_total: float | None
    cbs: float | None
    valor_total_nfse: float | None


class GeradorRelatoriosService:
    """Gera demonstrativos fiscais sem inventar valores ausentes nos XMLs."""

    AZUL = "1D4ED8"
    AZUL_ESCURO = "0F172A"

    def gerar(self, empresa: Empresa, periodo: PeriodoConsulta, preferencias: PreferenciasDownload) -> tuple[Path, ...]:
        if not preferencias.gerar_excel and not preferencias.gerar_pdf:
            return ()
        linhas = self._carregar_linhas(empresa, periodo, preferencias)
        pasta = OrganizadorDocumentosADN.pasta_da_empresa(empresa) / "Relatórios"
        pasta.mkdir(parents=True, exist_ok=True)
        sufixo = f"{periodo.data_inicio:%Y-%m-%d}_a_{periodo.data_fim:%Y-%m-%d}"
        caminhos: list[Path] = []
        if preferencias.gerar_excel:
            caminho = pasta / f"Relatorio_Fiscal_NFSe_{sufixo}.xlsx"
            self._gerar_excel(caminho, empresa, periodo, linhas)
            caminhos.append(caminho)
        if preferencias.gerar_pdf:
            caminho = pasta / f"Relatorio_Fiscal_NFSe_{sufixo}.pdf"
            self._gerar_pdf(caminho, empresa, periodo, linhas)
            caminhos.append(caminho)
        return tuple(caminhos)

    def _carregar_linhas(self, empresa, periodo, preferencias) -> tuple[LinhaRelatorio, ...]:
        arquivo = OrganizadorDocumentosADN.pasta_da_empresa(empresa) / OrganizadorDocumentosADN.PASTA_ARQUIVO_INTERNO
        if not arquivo.exists():
            return ()
        linhas: list[LinhaRelatorio] = []
        for caminho in arquivo.rglob("*.xml"):
            try:
                xml = caminho.read_bytes()
                if b"<!DOCTYPE" in xml[:2048].upper() or b"<!ENTITY" in xml[:2048].upper():
                    continue
                raiz = ElementTree.fromstring(xml)
                if self._local(raiz.tag) != "NFSe":
                    continue
                linha = self._extrair_linha(raiz, caminho)
            except (OSError, ElementTree.ParseError, ValueError):
                continue
            if not (periodo.data_inicio <= linha.data_emissao <= periodo.data_fim):
                continue
            if linha.direcao == DirecaoDocumento.EMITIDA.value and not preferencias.baixar_emitidas:
                continue
            if linha.direcao == DirecaoDocumento.RECEBIDA.value and not preferencias.baixar_recebidas:
                continue
            if linha.direcao != DirecaoDocumento.OUTRA.value:
                linhas.append(linha)
        return tuple(sorted(linhas, key=lambda x: (x.data_emissao, x.numero_nfse, x.chave)))

    def _extrair_linha(self, raiz, caminho) -> LinhaRelatorio:
        inf = self._elemento(raiz, "infNFSe")
        dps = self._elemento(inf, "infDPS")
        data_texto = self._texto(inf, "dhEmi") or self._texto(dps, "dhEmi")
        if not data_texto:
            raise ValueError("data ausente")
        prestador = self._grupo(dps, {"prest", "prestador"})
        tomador = self._grupo(dps, {"toma", "tomador"})
        intermediario = self._grupo(dps, {"interm", "intermediario"})
        servico = self._elemento(dps, "serv")
        cod_serv = self._elemento(servico, "cServ")
        valores_dps = self._filho(dps, "valores")
        valores_servico = self._filho(valores_dps, "vServPrest")
        trib = self._filho(valores_dps, "trib")
        trib_mun = self._filho(trib, "tribMun")
        trib_fed = self._filho(trib, "tribFed")
        pis_cofins = self._filho(trib_fed, "piscofins")
        total_trib = self._filho(trib, "totTrib")
        valores_nfse = self._filho(inf, "valores")
        ibscbs = self._filho(inf, "IBSCBS")
        valores_ibscbs = self._filho(ibscbs, "valores")
        total_ibscbs = self._filho(ibscbs, "totCIBS")
        gibs = self._filho(total_ibscbs, "gIBS")
        gibs_uf = self._filho(gibs, "gIBSUFTot")
        gibs_mun = self._filho(gibs, "gIBSMunTot")
        gcbs = self._filho(total_ibscbs, "gCBS")
        regime = self._elemento(dps, "regTrib")
        chave = str(inf.attrib.get("Id", "")) if inf is not None else ""
        chave = chave or caminho.stem.split("_", 1)[-1]
        direcao = caminho.parent.name
        if direcao not in {x.value for x in DirecaoDocumento}:
            direcao = DirecaoDocumento.OUTRA.value
        ibs_uf = self._numero(self._filho_texto(gibs_uf, "vIBSUF"))
        ibs_mun = self._numero(self._filho_texto(gibs_mun, "vIBSMun"))
        return LinhaRelatorio(
            self._data(data_texto) or date.min, self._data(self._texto(dps, "dCompet")), direcao,
            self._texto(inf, "nNFSe") or "", self._texto(dps, "serie") or "", self._texto(dps, "nDPS") or "", chave,
            self._texto(inf, "xLocEmi") or "", self._texto(inf, "xLocIncid") or "",
            self._filho_texto(cod_serv, "cTribNac") or "",
            self._filho_texto(cod_serv, "xDescServ") or self._texto(inf, "xTribNac") or "",
            self._filho_texto(cod_serv, "cNBS") or "", self._texto(inf, "xNBS") or "",
            prestador[0], prestador[1], tomador[0], tomador[1], intermediario[0], intermediario[1],
            self._filho_texto(regime, "opSimpNac") or "", self._filho_texto(regime, "regEspTrib") or "",
            self._filho_texto(trib_mun, "tpRetISSQN") or "",
            self._numero(self._filho_texto(valores_servico, "vServ")),
            self._numero(self._filho_texto(valores_dps, "vDescCondIncond")),
            self._numero(self._filho_texto(valores_nfse, "vBC")),
            self._numero(self._filho_texto(valores_nfse, "pAliqAplic")),
            self._numero(self._filho_texto(valores_nfse, "vISSQN")),
            self._numero(self._filho_texto(pis_cofins, "vPis")), self._numero(self._filho_texto(pis_cofins, "vCofins")),
            self._numero(self._filho_texto(trib_fed, "vRetIRRF")), self._numero(self._filho_texto(trib_fed, "vRetCSLL")),
            self._numero(self._filho_texto(trib_fed, "vRetCP") or self._filho_texto(trib_fed, "vRetINSS")),
            self._numero(self._filho_texto(valores_nfse, "vTotalRet")), self._numero(self._filho_texto(valores_nfse, "vLiq")),
            self._numero(self._filho_texto(total_trib, "vTotTribFed")), self._numero(self._filho_texto(total_trib, "vTotTribEst")),
            self._numero(self._filho_texto(total_trib, "vTotTribMun")), self._numero(self._filho_texto(total_trib, "vTotTrib")),
            self._numero(self._filho_texto(valores_ibscbs, "vBC")), ibs_uf, ibs_mun,
            (ibs_uf + ibs_mun) if ibs_uf is not None and ibs_mun is not None else None,
            self._numero(self._filho_texto(gcbs, "vCBS")), self._numero(self._filho_texto(total_ibscbs, "vTotNF")),
        )

    def _gerar_excel(self, destino, empresa, periodo, linhas) -> None:
        wb = Workbook()
        wb.remove(wb.active)
        grupos = (
            ("Emitidas", tuple(x for x in linhas if x.direcao == DirecaoDocumento.EMITIDA.value)),
            ("Recebidas", tuple(x for x in linhas if x.direcao == DirecaoDocumento.RECEBIDA.value)),
        )
        for categoria, grupo in grupos:
            if not grupo:
                continue
            resumo = wb.create_sheet(f"Resumo {categoria.lower()}")
            detalhe = wb.create_sheet(f"Notas {categoria.lower()}")
            self._montar_resumo_excel(resumo, empresa, periodo, grupo, categoria)
            self._montar_detalhe_excel(detalhe, grupo, f"TabelaNFSe{categoria}")
        if not wb.sheetnames:
            resumo = wb.create_sheet("Resumo fiscal")
            self._montar_resumo_excel(resumo, empresa, periodo, (), "Notas")
        self._salvar_atomico(destino, lambda temporario: wb.save(temporario))

    def _montar_resumo_excel(self, ws, empresa, periodo, linhas, categoria) -> None:
        ws.sheet_view.showGridLines = False
        ws["A2"] = f"Resumo fiscal - Notas {categoria.lower()}"
        ws["A2"].font = Font(name="Arial", size=16, bold=True, color=self.AZUL_ESCURO)
        ws["A3"] = empresa.nome_exibicao
        ws["A4"] = f"CNPJ: {empresa.cnpj_formatado}"
        ws["A5"] = f"Período: {periodo.data_inicio:%d/%m/%Y} a {periodo.data_fim:%d/%m/%Y}"
        ws["A7"] = "Visão geral"
        ws.append([])
        ws.append(["Indicador", "Total"])
        indicadores = [("Quantidade de notas", None), ("Valor dos serviços", "valor_servico"), ("Valor líquido", "valor_liquido"),
                       ("Total retido", "total_retido"), ("ISS", "valor_iss"), ("PIS", "valor_pis"),
                       ("COFINS", "valor_cofins"), ("IRRF", "valor_irrf"), ("CSLL", "valor_csll"),
                       ("INSS", "valor_inss"), ("IBS estadual", "ibs_uf"), ("IBS municipal", "ibs_municipal"), ("CBS", "cbs")]
        for rotulo, campo in indicadores:
            if campo is None:
                ws.append([rotulo, len(linhas)])
            else:
                ws.append([rotulo, self._soma(linhas, campo)])
        linha_secao = ws.max_row + 2
        ws.cell(linha_secao, 1, "Consolidação dos tributos informados")
        ws.append(["Tributo / valor", "Total", "Notas com informação"])
        tributos = [("Base de cálculo do ISS", "base_iss"), ("ISS", "valor_iss"), ("PIS", "valor_pis"),
                    ("COFINS", "valor_cofins"), ("IRRF", "valor_irrf"), ("CSLL", "valor_csll"), ("INSS", "valor_inss"),
                    ("Total retido informado", "total_retido"), ("Tributos federais aproximados", "tributos_federais"),
                    ("Tributos estaduais aproximados", "tributos_estaduais"), ("Tributos municipais aproximados", "tributos_municipais"),
                    ("Base IBS/CBS", "base_ibs_cbs"), ("IBS estadual", "ibs_uf"), ("IBS municipal", "ibs_municipal"), ("CBS", "cbs")]
        for rotulo, campo in tributos:
            ws.append([rotulo, self._soma(linhas, campo), sum(getattr(x, campo) is not None for x in linhas)])
        linha_contato = ws.max_row + 2
        contato = ws.cell(linha_contato, 1, f"Contato e suporte: Instagram {SUPPORT_HANDLE}")
        contato.hyperlink = SUPPORT_URL
        contato.style = "Hyperlink"
        contato.font = Font(name="Arial", size=10, bold=True, color=self.AZUL, underline="single")
        self._estilo_cabecalho(ws[9])
        self._estilo_cabecalho(ws[linha_secao + 1])
        for celula in (ws["A7"], ws.cell(linha_secao, 1)):
            celula.font = Font(name="Arial", size=12, bold=True, color=self.AZUL_ESCURO)
        for row in ws.iter_rows(min_row=10, max_row=ws.max_row, min_col=2, max_col=2):
            for cell in row:
                cell.number_format = 'R$ #,##0.00;[Red]-R$ #,##0.00;"-"'
        ws["B10"].number_format = "0"
        for row in ws.iter_rows(min_row=linha_secao + 2, max_row=ws.max_row, min_col=3, max_col=3):
            row[0].number_format = "0"
        ws.column_dimensions["A"].width = 40
        for col in "BC": ws.column_dimensions[col].width = 22
        ws.freeze_panes = "A9"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_area = f"A1:C{ws.max_row}"

    def _montar_detalhe_excel(self, ws, linhas, nome_tabela="TabelaFiscalNFSe") -> None:
        ws.sheet_view.showGridLines = False
        headers = ["Data", "Competência", "Tipo", "Nº NFS-e", "Série DPS", "Nº DPS", "Chave", "Município emissão",
                   "Município incidência", "Código serviço", "Descrição serviço", "Código NBS", "Descrição NBS", "CNPJ prestador",
                   "Prestador", "CNPJ tomador", "Tomador", "CNPJ intermediário", "Intermediário", "Opção Simples Nacional",
                   "Regime especial", "Retenção ISS", "Valor serviços", "Descontos", "Base ISS", "Alíquota ISS (%)", "ISS", "PIS",
                   "COFINS", "IRRF", "CSLL", "INSS", "Total retido", "Valor líquido", "Trib. federais", "Trib. estaduais",
                   "Trib. municipais", "Total tributos", "Base IBS/CBS", "IBS UF", "IBS Município", "IBS Total", "CBS", "Valor total NFS-e"]
        ws.append(headers)
        for x in linhas:
            ws.append([x.data_emissao, x.competencia, x.direcao, x.numero_nfse, x.serie_dps, x.numero_dps, x.chave,
                       x.municipio_emissao, x.municipio_incidencia, x.codigo_servico, x.descricao_servico, x.codigo_nbs, x.descricao_nbs,
                       x.prestador_cnpj, x.prestador_nome, x.tomador_cnpj, x.tomador_nome, x.intermediario_cnpj, x.intermediario_nome,
                       x.simples_nacional, x.regime_especial, x.tipo_retencao_iss, x.valor_servico, x.desconto, x.base_iss, x.aliquota_iss,
                       x.valor_iss, x.valor_pis, x.valor_cofins, x.valor_irrf, x.valor_csll, x.valor_inss, x.total_retido, x.valor_liquido,
                       x.tributos_federais, x.tributos_estaduais, x.tributos_municipais, x.total_tributos, x.base_ibs_cbs, x.ibs_uf,
                       x.ibs_municipal, x.ibs_total, x.cbs, x.valor_total_nfse])
        self._estilo_cabecalho(ws[1])
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 32
        for col in range(1, 45): ws.column_dimensions[ws.cell(1, col).column_letter].width = 18
        for col in (7, 11, 13, 15, 17, 19): ws.column_dimensions[ws.cell(1, col).column_letter].width = 35
        for row in ws.iter_rows(min_row=2):
            row[0].number_format = "dd/mm/yyyy"; row[1].number_format = "dd/mm/yyyy"
            for cell in row[22:]: cell.number_format = 'R$ #,##0.00;[Red]-R$ #,##0.00;""'
            row[25].number_format = "0.0000"
        if linhas:
            tabela = ExcelTable(displayName=nome_tabela, ref=f"A1:AR{len(linhas)+1}")
            tabela.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False, showColumnStripes=False)
            ws.add_table(tabela)
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:1"
        ws.oddFooter.center.text = f"Contato e suporte: Instagram {SUPPORT_HANDLE}"
        ws.oddFooter.center.size = 8
        ws.oddFooter.center.color = "64748B"

    def _gerar_pdf(self, destino, empresa, periodo, linhas) -> None:
        def construir(temporario):
            doc = SimpleDocTemplate(temporario, pagesize=landscape(A4), rightMargin=12*mm, leftMargin=12*mm, topMargin=12*mm,
                                    bottomMargin=14*mm, title="Relatório fiscal de NFS-e", author="NFS-e Fácil")
            s = self._estilos_pdf()
            emitidas = tuple(x for x in linhas if x.direcao == DirecaoDocumento.EMITIDA.value)
            recebidas = tuple(x for x in linhas if x.direcao == DirecaoDocumento.RECEBIDA.value)
            secoes = [
                ("NOTAS EMITIDAS", "Resumo Fiscal - Notas Emitidas", emitidas),
                ("NOTAS RECEBIDAS", "Resumo Fiscal - Notas Recebidas", recebidas),
            ]
            story = []
            for etiqueta, titulo, grupo in secoes:
                if not grupo:
                    continue
                if story:
                    story.append(PageBreak())
                story.extend([
                    *self._cabecalho_pdf(empresa, periodo, etiqueta, titulo=titulo),
                    self._cards_pdf(grupo, quantidade_rotulo=etiqueta),
                    Spacer(1, 5*mm),
                    Paragraph("Tributos, bases e retenções", s["secao"]),
                    self._tabela_tributos_pdf(grupo),
                    Spacer(1, 4*mm),
                    Paragraph(
                        "Valores calculados exclusivamente a partir dos XMLs oficiais desta categoria. Campos ausentes não são tratados como valores informados.",
                        s["nota"],
                    ),
                    PageBreak(),
                    *self._cabecalho_pdf(
                        empresa,
                        periodo,
                        f"DETALHAMENTO - {etiqueta}",
                        compacto=True,
                        titulo=f"Detalhamento - {etiqueta.title()}",
                    ),
                    self._tabela_notas_pdf(grupo),
                ])
            if not story:
                story = [
                    *self._cabecalho_pdf(empresa, periodo, "RELATÓRIO FISCAL", titulo="Relatório fiscal de NFS-e"),
                    self._cards_pdf(()),
                    Spacer(1, 5*mm),
                    Paragraph("Nenhuma nota foi localizada para os filtros selecionados.", s["normal"]),
                ]
            doc.build(story, onFirstPage=self._rodape_pdf, onLaterPages=self._rodape_pdf)
        self._salvar_atomico(destino, construir)

    def _cabecalho_pdf(self, empresa, periodo, etiqueta, compacto=False, titulo=None):
        s = self._estilos_pdf()
        titulo = titulo or ("Relatório fiscal de NFS-e" if not compacto else etiqueta.title())
        return [
            PDFTable(
                [[Paragraph(etiqueta, s["etiqueta"]), Paragraph("NFS-e Fácil", s["marca"])]],
                colWidths=[220*mm, 53*mm],
                style=TableStyle([("VALIGN", (0,0), (-1,-1), "MIDDLE"), ("ALIGN", (1,0), (1,0), "RIGHT"),
                                  ("BOTTOMPADDING", (0,0), (-1,-1), 2), ("TOPPADDING", (0,0), (-1,-1), 0)]),
            ),
            Paragraph(titulo, s["titulo"]),
            Paragraph(empresa.nome_exibicao, s["subtitulo"]),
            Paragraph(
                f"CNPJ {empresa.cnpj_formatado} &nbsp;&nbsp;&nbsp; Período {periodo.data_inicio:%d/%m/%Y} a {periodo.data_fim:%d/%m/%Y}",
                s["normal"],
            ),
            Spacer(1, 2.5*mm),
            HRFlowable(width="100%", thickness=1.2, color=colors.HexColor("#2563EB"), spaceAfter=5*mm),
        ]

    def _cards_pdf(self, linhas, quantidade_rotulo="NOTAS NO PERÍODO"):
        cards = [
            (quantidade_rotulo, str(len(linhas))),
            ("VALOR DOS SERVIÇOS", self._moeda(self._soma(linhas, "valor_servico"))),
            ("VALOR LÍQUIDO", self._moeda(self._soma(linhas, "valor_liquido"))),
            ("TOTAL RETIDO", self._moeda(self._soma(linhas, "total_retido"))),
        ]
        s = self._estilos_pdf()
        celulas = []
        for rotulo, valor in cards:
            celulas.append([Paragraph(rotulo, s["card_rotulo"]), Paragraph(valor, s["card_valor"])])
        tabela = PDFTable([celulas], colWidths=[66.75*mm]*4, hAlign="LEFT")
        estilo = [("VALIGN",(0,0),(-1,-1),"MIDDLE"),("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#F1F5F9")),
                  ("BOX",(0,0),(-1,-1),.6,colors.HexColor("#CBD5E1")),("INNERGRID",(0,0),(-1,-1),3,colors.white),
                  ("LEFTPADDING",(0,0),(-1,-1),10),("RIGHTPADDING",(0,0),(-1,-1),10),
                  ("TOPPADDING",(0,0),(-1,-1),8),("BOTTOMPADDING",(0,0),(-1,-1),8)]
        tabela.setStyle(TableStyle(estilo))
        return tabela

    def _tabela_resumo_direcao(self, linhas, incluir_principais=True):
        emitidas = tuple(x for x in linhas if x.direcao == DirecaoDocumento.EMITIDA.value)
        recebidas = tuple(x for x in linhas if x.direcao == DirecaoDocumento.RECEBIDA.value)
        dados = [["Indicador", "Emitidas", "Recebidas", "Total"]]
        indicadores = (("Quantidade de notas", None), ("Valor dos serviços", "valor_servico"), ("Valor líquido", "valor_liquido"),
                       ("Total retido", "total_retido"), ("ISS", "valor_iss"), ("IBS total", "ibs_total"), ("CBS", "cbs"))
        if not incluir_principais:
            indicadores = (("Quantidade de notas", None), ("ISS", "valor_iss"), ("IBS total", "ibs_total"), ("CBS", "cbs"))
        for rotulo, campo in indicadores:
            vals = (len(emitidas), len(recebidas), len(linhas)) if campo is None else (self._soma(emitidas,campo), self._soma(recebidas,campo), self._soma(linhas,campo))
            dados.append([rotulo, *([str(v) for v in vals] if campo is None else [self._moeda(v) for v in vals])])
        tabela = self._tabela_pdf(dados, [75*mm, 66*mm, 66*mm, 66*mm])
        tabela.setStyle(TableStyle([
            ("ALIGN", (0, 0), (0, -1), "LEFT"),
            ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ]))
        return tabela

    def _tabela_tributos_pdf(self, linhas):
        itens = []
        for rotulo, campo in (("Base de cálculo do ISS","base_iss"),("ISS","valor_iss"),("PIS","valor_pis"),("COFINS","valor_cofins"),
                              ("IRRF","valor_irrf"),("CSLL","valor_csll"),("INSS","valor_inss"),("Total retido informado","total_retido"),
                              ("Base IBS/CBS","base_ibs_cbs"),("IBS estadual","ibs_uf"),("IBS municipal","ibs_municipal"),("CBS","cbs")):
            qtd = sum(getattr(x,campo) is not None for x in linhas)
            itens.append([rotulo, self._moeda(self._soma(linhas,campo)) if qtd else "Não informado", qtd])
        metade = (len(itens) + 1) // 2
        grupos = (itens[:metade], itens[metade:])
        tabelas = []
        for grupo in grupos:
            dados = [["Tributo / base", "Total", "Notas"], *grupo]
            tabela = self._tabela_pdf(dados, [62*mm, 48*mm, 22.5*mm])
            tabela.setStyle(TableStyle([
                ("ALIGN", (0, 0), (0, -1), "LEFT"),
                ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
            ]))
            tabelas.append(tabela)

        # Duas tabelas independentes preservam o eixo de cada coluna e criam
        # uma separação visual real entre os dois blocos fiscais.
        return PDFTable(
            [[tabelas[0], "", tabelas[1]]],
            colWidths=[132.5*mm, 8*mm, 132.5*mm],
            hAlign="LEFT",
            style=TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
            ]),
        )

    def _tabela_notas_pdf(self, linhas):
        s = self._estilos_pdf()["tabela"]
        dados = [["Data","Nº NFS-e","Contraparte","Serviço","Vlr. serviço","Base ISS","ISS","Retido","Líquido"]]
        for x in linhas:
            contraparte = (x.tomador_nome or x.tomador_cnpj) if x.direcao == DirecaoDocumento.EMITIDA.value else (x.prestador_nome or x.prestador_cnpj)
            dados.append([x.data_emissao.strftime("%d/%m/%Y"),x.numero_nfse or "-",Paragraph(self._resumir(contraparte,34),s),
                          Paragraph(self._resumir(x.descricao_servico or x.codigo_servico,34),s),self._moeda(x.valor_servico),self._moeda(x.base_iss),
                          self._moeda(x.valor_iss),self._moeda(x.total_retido),self._moeda(x.valor_liquido)])
        tabela = self._tabela_pdf(dados, [20*mm,28*mm,50*mm,61*mm,25*mm,24*mm,20*mm,20*mm,25*mm])
        tabela.setStyle(TableStyle([
            ("ALIGN", (0, 1), (1, -1), "CENTER"),
            ("ALIGN", (2, 1), (3, -1), "LEFT"),
            ("ALIGN", (4, 1), (8, -1), "RIGHT"),
        ]))
        return tabela

    @staticmethod
    def _tabela_pdf(dados, larguras):
        t = PDFTable(dados, repeatRows=1, colWidths=larguras, hAlign="LEFT")
        t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#0F2A5F")),("TEXTCOLOR",(0,0),(-1,0),colors.white),
                               ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("FONTNAME",(0,1),(-1,-1),"Helvetica"),("FONTSIZE",(0,0),(-1,-1),7.2),
                               ("ALIGN",(0,0),(-1,0),"CENTER"),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("ALIGN",(1,1),(-1,-1),"RIGHT"),
                               ("ALIGN",(0,1),(0,-1),"LEFT"),
                               ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F8FAFC")]),
                               ("LINEBELOW",(0,0),(-1,-1),.25,colors.HexColor("#CBD5E1")),("TOPPADDING",(0,0),(-1,-1),3.2),("BOTTOMPADDING",(0,0),(-1,-1),3.2)]))
        return t

    @staticmethod
    def _estilos_pdf():
        b=getSampleStyleSheet()
        return {"titulo":ParagraphStyle("Titulo",parent=b["Title"],fontName="Helvetica-Bold",fontSize=17,leading=21,textColor=colors.HexColor("#0F172A"),alignment=TA_LEFT,spaceAfter=4),
                "subtitulo":ParagraphStyle("Subtitulo",parent=b["Normal"],fontName="Helvetica-Bold",fontSize=10,leading=13,textColor=colors.HexColor("#334155")),
                "normal":ParagraphStyle("NormalFiscal",parent=b["Normal"],fontName="Helvetica",fontSize=9,leading=12,textColor=colors.HexColor("#334155")),
                "secao":ParagraphStyle("Secao",parent=b["Heading2"],fontName="Helvetica-Bold",fontSize=11,leading=14,textColor=colors.HexColor("#0F172A"),spaceAfter=5),
                "nota":ParagraphStyle("Nota",parent=b["Normal"],fontName="Helvetica-Oblique",fontSize=7.5,leading=10,textColor=colors.HexColor("#64748B")),
                "tabela":ParagraphStyle("Tabela",parent=b["Normal"],fontName="Helvetica",fontSize=7.2,leading=8.6,textColor=colors.HexColor("#1F2937")),
                "etiqueta":ParagraphStyle("Etiqueta",parent=b["Normal"],fontName="Helvetica-Bold",fontSize=7.5,leading=9,textColor=colors.HexColor("#2563EB"),spaceAfter=2),
                "marca":ParagraphStyle("Marca",parent=b["Normal"],fontName="Helvetica-Bold",fontSize=9,leading=11,textColor=colors.HexColor("#0F2A5F"),alignment=2),
                "card_rotulo":ParagraphStyle("CardRotulo",parent=b["Normal"],fontName="Helvetica-Bold",fontSize=6.8,leading=8,textColor=colors.HexColor("#64748B"),spaceAfter=3),
                "card_valor":ParagraphStyle("CardValor",parent=b["Normal"],fontName="Helvetica-Bold",fontSize=13,leading=15,textColor=colors.HexColor("#0F172A"))}

    @staticmethod
    def _rodape_pdf(canvas, doc):
        canvas.saveState(); canvas.setFont("Helvetica",7); canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawString(12*mm,7*mm,"NFS-e Fácil - Relatório gerado a partir dos XMLs oficiais")
        contato = f"Contato: Instagram {SUPPORT_HANDLE}"
        largura = canvas.stringWidth(contato, "Helvetica", 7)
        x_contato = (landscape(A4)[0] - largura) / 2
        canvas.setFillColor(colors.HexColor("#2563EB"))
        canvas.drawString(x_contato, 7*mm, contato)
        canvas.linkURL(
            SUPPORT_URL,
            (x_contato, 6.2*mm, x_contato + largura, 9*mm),
            relative=0,
            thickness=0,
        )
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawRightString(landscape(A4)[0]-12*mm,7*mm,f"Página {doc.page}"); canvas.restoreState()

    def _estilo_cabecalho(self, celulas):
        for c in celulas:
            c.fill=PatternFill("solid",fgColor=self.AZUL); c.font=Font(name="Arial",size=10,bold=True,color="FFFFFF")
            c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)

    @staticmethod
    def _soma(linhas,campo): return sum(v for x in linhas if (v:=getattr(x,campo)) is not None)
    @staticmethod
    def _moeda(valor):
        if valor is None: return "Não informado"
        return f"R$ {valor:,.2f}".replace(",","X").replace(".",",").replace("X",".")
    @staticmethod
    def _data(valor):
        if not valor: return None
        try: return date.fromisoformat(valor[:10])
        except ValueError: return None
    def _grupo(self,raiz,nomes):
        if raiz is None: return "",""
        for e in raiz.iter():
            if self._local(e.tag).casefold() in nomes:
                return self._texto(e,"CNPJ") or self._texto(e,"CPF") or "", self._texto(e,"xNome") or self._texto(e,"Nome") or ""
        return "",""
    @staticmethod
    def _local(tag): return tag.rsplit("}",1)[-1]
    @classmethod
    def _elemento(cls,raiz,nome):
        if raiz is None: return None
        return next((e for e in raiz.iter() if cls._local(e.tag)==nome),None)
    @classmethod
    def _filho(cls,raiz,nome):
        if raiz is None: return None
        return next((e for e in raiz if cls._local(e.tag)==nome),None)
    @classmethod
    def _texto(cls,raiz,nome):
        e=cls._elemento(raiz,nome); return e.text.strip() if e is not None and e.text else None
    @classmethod
    def _filho_texto(cls,raiz,nome):
        e=cls._filho(raiz,nome); return e.text.strip() if e is not None and e.text else None
    @staticmethod
    def _numero(valor):
        if not valor: return None
        try: return float(valor.replace(",","."))
        except ValueError: return None
    @staticmethod
    def _resumir(texto,limite):
        texto=re.sub(r"\s+"," ",texto).strip(); return texto if len(texto)<=limite else texto[:limite-3]+"..."
    @staticmethod
    def _salvar_atomico(destino,escritor):
        fd,tmp=tempfile.mkstemp(prefix=".relatorio_",suffix=destino.suffix,dir=destino.parent); os.close(fd)
        try: escritor(tmp); os.replace(tmp,destino)
        except Exception as err: raise DocumentoFiscalGravacaoError("Não foi possível gerar o relatório solicitado.") from err
        finally:
            if os.path.exists(tmp): os.unlink(tmp)


RelatorioService = GeradorRelatoriosService
